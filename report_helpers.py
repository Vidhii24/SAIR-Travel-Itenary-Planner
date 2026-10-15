import os
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=140, bottom=140, left=180, right=180):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="D1D5DB", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def add_callout(doc, text_content, title="NOTE / ARCHITECTURAL PRINCIPLE", color_hex="264638", bg_hex="F4F6F4"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    cell = tbl.rows[0].cells[0]
    cell.width = Inches(6.5)
    set_cell_background(cell, bg_hex)
    set_cell_margins(cell, top=160, bottom=160, left=240, right=200)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:left w:val="single" w:sz="36" w:space="0" w:color="{color_hex}"/>'
        f'<w:top w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:bottom w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    run_title = p.add_run(f"📌 {title}\n")
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(10)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0x26, 0x46, 0x38)
    
    run_body = p.add_run(text_content)
    run_body.font.name = "Calibri"
    run_body.font.size = Pt(9.5)
    run_body.font.italic = True
    run_body.font.color.rgb = RGBColor(0x37, 0x41, 0x51)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

def add_figure(doc, img_path, fig_num, title, description):
    if not os.path.exists(img_path):
        print(f"Warning: image {img_path} not found!")
        return
        
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(12)
    p_img.paragraph_format.space_after = Pt(4)
    run_img = p_img.add_run()
    run_img.add_picture(img_path, width=Inches(6.0))
    
    p_caption = doc.add_paragraph()
    p_caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_caption.paragraph_format.space_before = Pt(2)
    p_caption.paragraph_format.space_after = Pt(4)
    run_caption = p_caption.add_run(f"Figure {fig_num}: {title}")
    run_caption.font.name = "Calibri"
    run_caption.font.size = Pt(10)
    run_caption.font.bold = True
    run_caption.font.color.rgb = RGBColor(0x26, 0x46, 0x38)
    
    p_desc = doc.add_paragraph()
    p_desc.paragraph_format.space_before = Pt(2)
    p_desc.paragraph_format.space_after = Pt(10)
    p_desc.paragraph_format.line_spacing = 1.15
    run_desc_label = p_desc.add_run("Functional Description: ")
    run_desc_label.font.name = "Calibri"
    run_desc_label.font.size = Pt(9.5)
    run_desc_label.font.bold = True
    run_desc_label.font.color.rgb = RGBColor(0x4B, 0x55, 0x63)
    
    run_desc = p_desc.add_run(description)
    run_desc.font.name = "Calibri"
    run_desc.font.size = Pt(9.5)
    run_desc.font.color.rgb = RGBColor(0x37, 0x41, 0x51)

print("Report helper modules initialized.")
